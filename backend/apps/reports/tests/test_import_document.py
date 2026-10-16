"""从文档导入战报正文：docx / txt → Markdown，不落库、不建附件。"""

from __future__ import annotations

import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.games.models import GAME_RECRUITING
from apps.reports.models import BODY_MAX_LENGTH, Report, ReportAttachment

pytestmark = pytest.mark.django_db

URL = reverse("v1:reports:report-import-document")


def make_docx(build) -> bytes:
    from docx import Document

    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_upload(data: bytes, name: str = "战报.docx"):
    return SimpleUploadedFile(
        name,
        data,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def txt_upload(data: bytes, name: str = "战报.txt"):
    return SimpleUploadedFile(name, data, content_type="text/plain")


@pytest.fixture
def attendee(game_factory, attendee_factory):
    game = game_factory()
    return game, attendee_factory(game)


def post(api_client, game, upload):
    return api_client.post(URL, {"game": game.pk, "file": upload}, format="multipart")


def test_docx_标题段落列表粗体都转成Markdown(api_client, attendee):
    game, user = attendee
    api_client.force_authenticate(user)

    def build(document):
        document.add_heading("第一幕 雾港", level=1)
        document.add_heading("码头", level=2)
        paragraph = document.add_paragraph("调查员在")
        paragraph.add_run("午夜").bold = True
        paragraph.add_run("抵达，")
        paragraph.add_run("雾很浓").italic = True
        paragraph.add_run("。")
        document.add_paragraph("手电筒", style="List Bullet")
        document.add_paragraph("左轮手枪", style="List Bullet")
        document.add_paragraph("先检定侦查", style="List Number")
        document.add_paragraph("再检定聆听", style="List Number")

    response = post(api_client, game, docx_upload(make_docx(build)))

    assert response.status_code == 200, response.content
    data = response.json()
    body = data["body_markdown"]
    assert "# 第一幕 雾港" in body
    assert "## 码头" in body
    assert "调查员在**午夜**抵达，*雾很浓*。" in body
    assert "- 手电筒" in body and "- 左轮手枪" in body
    assert "1. 先检定侦查" in body and "2. 再检定聆听" in body
    assert data["truncated"] is False
    assert data["warnings"] == []


def test_导入不落库也不建附件(api_client, attendee):
    game, user = attendee
    api_client.force_authenticate(user)

    response = post(api_client, game, txt_upload("正文".encode()))

    assert response.status_code == 200
    assert not Report.objects.exists()
    assert not ReportAttachment.objects.exists()


@pytest.mark.parametrize(
    ("raw", "gbk_warning"),
    [
        ("雾港疑影\n第二段".encode(), False),
        (b"\xef\xbb\xbf" + "雾港疑影\n第二段".encode(), False),
        ("雾港疑影\n第二段".encode("gbk"), True),
    ],
    ids=["utf-8", "utf-8-sig", "gbk"],
)
def test_txt_三种编码都能读(api_client, attendee, raw, gbk_warning):
    game, user = attendee
    api_client.force_authenticate(user)

    response = post(api_client, game, txt_upload(raw))

    assert response.status_code == 200, response.content
    data = response.json()
    assert data["body_markdown"] == "雾港疑影\n\n第二段"  # BOM 不得残留
    assert any("GBK" in w for w in data["warnings"]) is gbk_warning


def test_txt_换行归一化且每行自成一段(api_client, attendee):
    """记事本写的战报通常一行一段；按 Markdown 语义单换行会把整篇并成一段，所以逐行分段。"""
    game, user = attendee
    api_client.force_authenticate(user)

    response = post(api_client, game, txt_upload("一\r\n二\r三\n\n\n\n四".encode()))

    assert response.json()["body_markdown"] == "一\n\n二\n\n三\n\n四"


def test_txt_里的Markdown表格不被拆开(api_client, attendee):
    game, user = attendee
    api_client.force_authenticate(user)
    raw = "前言\n| 属性 | 值 |\n| --- | --- |\n| 力量 | 60 |\n后记"

    response = post(api_client, game, txt_upload(raw.encode()))

    assert response.json()["body_markdown"] == "前言\n\n| 属性 | 值 |\n| --- | --- |\n| 力量 | 60 |\n\n后记"


def test_无法识别的编码回400(api_client, attendee):
    game, user = attendee
    api_client.force_authenticate(user)

    # 0x80 在 gbk 里是非法首字节，utf-8 也解不了
    response = post(api_client, game, txt_upload(b"\x80\x80\x80abc"))

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "import_unknown_encoding"
    assert "UTF-8" in response.json()["error"]["message"]


def test_超长正文截断并标记(api_client, attendee):
    game, user = attendee
    api_client.force_authenticate(user)
    line = "字" * 99 + "\n"
    raw = (line * 1000).encode()  # 100000 字符

    response = post(api_client, game, txt_upload(raw))

    data = response.json()
    assert response.status_code == 200
    assert data["truncated"] is True
    assert len(data["body_markdown"]) <= BODY_MAX_LENGTH
    assert len(data["body_markdown"]) > BODY_MAX_LENGTH - 2000
    assert data["body_markdown"].endswith("字")  # 断在行边界上


def test_未登录401(api_client, attendee):
    game, _ = attendee

    response = post(api_client, game, txt_upload(b"x"))

    assert response.status_code == 401


def test_不在团里的人403(api_client, attendee, user_factory):
    game, _ = attendee
    api_client.force_authenticate(user_factory())

    response = post(api_client, game, txt_upload(b"x"))

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "not_in_game"


def test_未结团400(api_client, game_factory, attendee_factory):
    game = game_factory(status=GAME_RECRUITING)
    user = attendee_factory(game)
    api_client.force_authenticate(user)

    response = post(api_client, game, txt_upload(b"x"))

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "game_not_finished"


def test_本团GM可以导入(api_client, game_factory):
    game = game_factory()
    api_client.force_authenticate(game.gm)

    assert post(api_client, game, txt_upload(b"x")).status_code == 200


def test_已有战报的作者可以导入(api_client, game_factory, user_factory, report_factory):
    """作者改稿走 `can_edit`（只看作者），导入与之同口径。"""
    game = game_factory()
    author = user_factory()
    report_factory(game, author)
    api_client.force_authenticate(author)

    assert post(api_client, game, txt_upload(b"x")).status_code == 200


def test_封禁号403(api_client, attendee):
    game, user = attendee
    user.is_banned = True
    user.save(update_fields=["is_banned"])
    api_client.force_authenticate(user)

    response = post(api_client, game, txt_upload(b"x"))

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "account_banned"


def test_缺文件或缺团400(api_client, attendee):
    game, user = attendee
    api_client.force_authenticate(user)

    assert api_client.post(URL, {"game": game.pk}, format="multipart").status_code == 400
    assert api_client.post(URL, {"file": txt_upload(b"x")}, format="multipart").status_code == 400


def test_docx_图片丢弃并提示(api_client, attendee):
    from docx.shared import Inches
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (4, 4), "red").save(buffer, format="PNG")
    png = buffer.getvalue()
    game, user = attendee
    api_client.force_authenticate(user)

    def build(document):
        document.add_paragraph("图前")
        document.add_picture(io.BytesIO(png), width=Inches(1))
        document.add_paragraph("图后")

    response = post(api_client, game, docx_upload(make_docx(build)))

    data = response.json()
    assert response.status_code == 200, response.content
    assert "图前" in data["body_markdown"] and "图后" in data["body_markdown"]
    assert "data:" not in data["body_markdown"] and "![" not in data["body_markdown"]
    assert any("图片" in w for w in data["warnings"])


def test_docx_表格转成表格(api_client, attendee):
    game, user = attendee
    api_client.force_authenticate(user)

    def build(document):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "属性"
        table.cell(0, 1).text = "值"
        table.cell(1, 0).text = "力量"
        table.cell(1, 1).text = "60"

    response = post(api_client, game, docx_upload(make_docx(build)))

    body = response.json()["body_markdown"]
    assert "| 属性 | 值 |" in body
    assert "| 力量 | 60 |" in body
    assert any("表格" in w for w in response.json()["warnings"])
